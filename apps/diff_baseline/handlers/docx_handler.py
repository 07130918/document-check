"""
DOCX Handler using python-docx
"""
from typing import List
from io import BytesIO
from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_COLOR_INDEX
from ..models import BBoxTextData


class DocxService:
    """python-docxを使用したWord文書処理サービス"""
    
    def extract_bbox_data(self, docx_bytes: bytes) -> List[BBoxTextData]:
        """
        Word文書からbbox_text_dataを抽出
        注: Wordでは直接的なbbox情報が取得できないため、擬似的に生成
        
        Args:
            docx_bytes: Word文書のバイトデータ
            
        Returns:
            bbox_text_dataのリスト
        """
        bbox_data_list = []
        
        # Word文書を開く
        doc = Document(BytesIO(docx_bytes))
        
        # ページ番号（Wordでは正確なページ番号が取得できないため推定）
        page_num = 0
        y_position = 0
        page_height = 792  # Letter sizeの高さ（ポイント）
        line_height = 20
        margin_left = 72  # 左マージン（1インチ）
        
        for paragraph in doc.paragraphs:
            # 段落内のランを処理
            x_position = margin_left
            
            # 段落のテキストを単語に分割
            words = paragraph.text.split()
            word_index = 0
            
            for run in paragraph.runs:
                run_words = run.text.split()
                
                for word in run_words:
                    if not word:
                        continue
                    
                    # 単語の幅を推定（文字数 × 平均文字幅）
                    word_width = len(word) * 10
                    
                    # 改ページの推定
                    if y_position > page_height - 72:  # 下マージンを考慮
                        page_num += 1
                        y_position = 72  # 上マージン
                    
                    bbox_data = BBoxTextData(
                        bbox=[x_position, y_position, word_width, line_height],
                        text=word,
                        page=page_num
                    )
                    bbox_data_list.append(bbox_data)
                    
                    # 次の単語の位置
                    x_position += word_width + 5  # スペース分
                    word_index += 1
            
            # 次の段落の位置
            y_position += line_height * 1.5  # 段落間隔
            
        return bbox_data_list
    
    def add_highlights_to_docx(self, docx_bytes: bytes, highlights: List[dict]) -> bytes:
        """
        Word文書にハイライトを追加
        
        Args:
            docx_bytes: 元のWord文書のバイトデータ
            highlights: ハイライト情報のリスト
            
        Returns:
            ハイライトが追加されたWord文書のバイトデータ
        """
        doc = Document(BytesIO(docx_bytes))
        
        # ハイライト情報を単語と照合
        highlight_map = {}
        for highlight in highlights:
            # 簡易的な実装：テキストベースでマッチング
            # 実際にはより高度なマッチングアルゴリズムが必要
            text = highlight.get('text', '')
            if text:
                highlight_map[text] = highlight['color']
        
        # 段落を処理
        for paragraph in doc.paragraphs:
            new_runs = []
            for run in paragraph.runs:
                words = run.text.split()
                for i, word in enumerate(words):
                    if word in highlight_map:
                        # 新しいランを作成してハイライトを適用
                        new_run = paragraph.add_run(word)
                        new_run.font.highlight_color = self._get_highlight_color(highlight_map[word])
                        if i < len(words) - 1:
                            paragraph.add_run(' ')
                    else:
                        # 通常のテキスト
                        new_run = paragraph.add_run(word)
                        if i < len(words) - 1:
                            paragraph.add_run(' ')
                
                # 元のランをクリア
                run.clear()
        
        # バイトデータとして保存
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        return output.read()
    
    def _get_highlight_color(self, color_name: str):
        """色名をWord のハイライトカラーに変換"""
        colors = {
            'red': WD_COLOR_INDEX.RED,
            'green': WD_COLOR_INDEX.BRIGHT_GREEN,
            'yellow': WD_COLOR_INDEX.YELLOW,
            'blue': WD_COLOR_INDEX.CYAN
        }
        return colors.get(color_name, WD_COLOR_INDEX.YELLOW)